"""문서 md를 Word(.docx)·PDF 보고서로 바꾼다 (`[1] docs/4) workflow/3_doc_rules.md`).

먼저 md_check로 검사하고, 오류가 없을 때만 md 옆에 같은 이름의 .docx(.pdf)를 만든다.
사용: python md_to_docx.py [--no-pdf] <md 파일>...
"""

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

import md_check

FONT = "Pretendard"
MONO = "Consolas"
TEXT = RGBColor.from_string("1A1A1A")
MUTED = RGBColor.from_string("6B6B6B")
BLUE = RGBColor.from_string("0072C6")
ADOPTED = RGBColor.from_string("15803D")
REJECTED = RGBColor.from_string("B91C1C")
CALLOUT_FILL, CALLOUT_BAR = "EAF6FF", "0099FF"
GRAY_FILL, TABLE_LINE = "F2F2F2", "D9D9D9"

FENCE = re.compile(r"^\s*```")
HEADING = re.compile(r"^(#+) (.*)$")
LIST_ITEM = re.compile(r"^( *)(-|(\d+)\.) (.*)$")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
TABLE_RULE = re.compile(r"^\|[\s:|-]+\|$")

# w:pPr 안에서 w:pBdr·w:shd 뒤에 와야 하는 요소들 (순서가 틀리면 Word가 못 연다)
PPR_AFTER_SHD = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
)


# ---------- 스타일 ----------

def _set_fonts(style) -> None:
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{key}"), None)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)


def _setup_styles(doc) -> None:
    normal = doc.styles["Normal"]
    _set_fonts(normal)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)

    title = doc.styles["Title"]
    _set_fonts(title)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = TEXT

    sizes = {"Heading 1": (18, 24, 12), "Heading 2": (14, 18, 6), "Heading 3": (12, 12, 4)}
    for name, (size, before, after) in sizes.items():
        style = doc.styles[name]
        _set_fonts(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = TEXT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True


def _setup_page(doc) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Mm(25))

    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field.append(OxmlElement("w:r"))
    p._p.append(field)


# ---------- 블록 모양 ----------

def _box(p, fill: str, bar: str | None = None) -> None:
    """문단에 배경색(과 왼쪽 막대)을 칠한다."""
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.insert_element_before(shd, *PPR_AFTER_SHD)
    if bar:
        border = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        for key, value in (("val", "single"), ("sz", "24"), ("space", "8"), ("color", bar)):
            left.set(qn(f"w:{key}"), value)
        border.append(left)
        ppr.insert_element_before(border, "w:shd", *PPR_AFTER_SHD)


def _inline(p, text: str) -> None:
    """**굵게**는 굵은 파랑, `코드`는 고정폭 글꼴로 넣는다."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
            run.font.color.rgb = BLUE
        elif part.startswith("`") and part.endswith("`"):
            p.add_run(part[1:-1]).font.name = MONO
        else:
            p.add_run(part)


def _paragraph(doc, text: str):
    p = doc.add_paragraph()
    if text.startswith(("**Q.**", "**A.**")):
        _inline(p, text)
        _box(p, GRAY_FILL)
    elif re.match(r"^[①-⑩]", text):
        run = p.add_run(text[0])
        run.bold = True
        run.font.color.rgb = BLUE
        _inline(p, text[1:])
    else:
        _inline(p, text)
    if text.startswith("근거:"):
        _muted(p)
    return p


def _muted(p) -> None:
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def _list_item(doc, indent: str, number: str | None, text: str, last: bool) -> None:
    level = 1 if len(indent) >= 4 else 0
    if number:
        mark = f"({number})" if level else f"{number}."
    else:
        mark = "–" if level else "•"
    p = doc.add_paragraph()
    p.add_run(f"{mark} ")
    _inline(p, text)
    if text.startswith("근거:"):
        _muted(p)
    p.paragraph_format.left_indent = Cm(1.2 if level else 0.5)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(6 if last else 2)


def _code_line(doc, line: str, last: bool) -> None:
    p = doc.add_paragraph()
    if line:
        run = p.add_run(line)
        run.font.name = MONO
        run.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(6 if last else 0)
    _box(p, GRAY_FILL)


def _cell_fill(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _table_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", TABLE_LINE)):
            el.set(qn(f"w:{key}"), value)
        borders.append(el)
    table._tbl.tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption"
    )


def _table(doc, lines: list[str]) -> None:
    rows = [
        [c.strip() for c in line.strip().strip("|").split("|")]
        for line in lines
        if not TABLE_RULE.match(line.strip())
    ]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    _table_borders(table)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = r < len(rows) - 1  # 표가 쪽 사이로 갈라지지 않게
            _inline(p, text)
            if r == 0:
                _cell_fill(cell, GRAY_FILL)
                for run in p.runs:
                    run.bold = True
            elif text.startswith(("채택", "기각")):
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = ADOPTED if text.startswith("채택") else REJECTED
    doc.add_paragraph()


def _toc(doc) -> None:
    head = doc.add_paragraph()
    head.paragraph_format.page_break_before = True
    run = head.add_run("목차")
    run.bold = True
    run.font.size = Pt(18)

    p = doc.add_paragraph()
    begin, instr, sep, end = (OxmlElement(t) for t in ("w:fldChar", "w:instrText", "w:fldChar", "w:fldChar"))
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    sep.set(qn("w:fldCharType"), "separate")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep):
        p.add_run()._r.append(el)
    p.add_run("PDF를 만들 때 목차가 채워집니다.")
    p.add_run()._r.append(end)


# ---------- 공개 함수 ----------

def build(text: str, toc: bool = False):
    doc = Document()
    _setup_styles(doc)
    _setup_page(doc)

    chapter = section = 0
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        i += 1
        if FENCE.match(line):
            while i < len(lines) and not FENCE.match(lines[i]):
                _code_line(doc, lines[i], last=bool(i + 1 < len(lines) and FENCE.match(lines[i + 1])))
                i += 1
            i += 1
            continue
        if not line.strip():
            continue

        heading = HEADING.match(line)
        if heading:
            depth, title = len(heading.group(1)), heading.group(2).strip()
            if depth == 1:
                doc.add_paragraph(title, style="Title")
                continue
            if depth == 2 and toc and chapter == 0:
                _toc(doc)
            p = doc.add_paragraph(style=f"Heading {depth - 1}")
            if depth == 2:
                chapter, section = chapter + 1, 0
                number = f"{chapter:02d}"
            elif depth == 3:
                section += 1
                number = f"{chapter}.{section}"
            else:
                number = None
            if number:
                p.add_run(f"{number}  ").font.color.rgb = BLUE
            p.add_run(title)
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            _inline(p, line[2:])
            _box(p, CALLOUT_FILL, CALLOUT_BAR)
            continue

        if line.startswith("|"):
            block = [line]
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            _table(doc, block)
            continue

        item = LIST_ITEM.match(line)
        if item:
            last = not (i < len(lines) and LIST_ITEM.match(lines[i]))
            _list_item(doc, item.group(1), item.group(3), item.group(4), last)
            continue

        _paragraph(doc, line)
    return doc


def to_pdf(docx_path: Path) -> Path:
    """로컬 Word로 목차를 채우고 PDF로 저장한다. Word가 있는 Windows에서만 된다."""
    pdf_path = docx_path.with_suffix(".pdf")

    def q(path: Path) -> str:
        return "'" + str(path.resolve()).replace("'", "''") + "'"

    script = f"""
$w = New-Object -ComObject Word.Application
$w.Visible = $false
try {{
    $d = $w.Documents.Open({q(docx_path)})
    foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
    $d.Save()
    $d.SaveAs2({q(pdf_path)}, 17)
    $d.Close()
}} finally {{
    $w.Quit()
}}
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    return pdf_path


def main(paths: list[str], pdf: bool = True) -> int:
    failed = False
    for p in paths:
        path = Path(p)
        if md_check.main([p]) != 0:
            failed = True
            print(f"{p}: 오류가 있어 Word를 만들지 않았습니다")
            continue
        text = path.read_text(encoding="utf-8")
        out = path.with_suffix(".docx")
        build(text, toc=path.stem.endswith("_full")).save(out)
        print(out)
        if pdf:
            print(to_pdf(out))
    return 1 if failed else 0


if __name__ == "__main__":
    args = sys.argv[1:]
    sys.exit(main([a for a in args if a != "--no-pdf"], pdf="--no-pdf" not in args))
