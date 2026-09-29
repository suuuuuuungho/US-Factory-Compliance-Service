"""SUU-297: 문서 md를 Word 보고서로 바꾼다 (`[1] docs/4) workflow/3_doc_rules.md`)."""

import re
import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from md_to_docx import build, main

PROJECT = Path(__file__).parents[3] / "1) project"

SAMPLE = """# 문서 제목

표지 문단.

## 첫 장

> 첫 장 핵심 문장입니다.

본문에 **강조**와 `applies` 코드가 있습니다.

### 첫 절

1. 가
    1. 가1
    2. 가2
2. 나

- 하나
    - 둘

근거: 40 CFR 70.6

### 둘째 절

#### 소절

| 방법 | 결과 |
|---|---|
| Kanon 2 | 채택 |
| BM25 | 기각 |

## 둘째 장

> 둘째 장 핵심 문장입니다.

### 셋째 절

① 첫째 포인트.

**Q.** 질문인가요?

**A.** 답입니다.

```text
원문 줄
  들여쓴 줄
```
"""


@pytest.fixture(scope="module")
def d():
    return build(SAMPLE)


def para(doc, text):
    """글자가 정확히 text인 문단 하나."""
    found = [p for p in doc.paragraphs if p.text == text]
    assert len(found) == 1, (text, [p.text for p in doc.paragraphs])
    return found[0]


def para_match(doc, pattern):
    found = [p for p in doc.paragraphs if re.fullmatch(pattern, p.text)]
    assert len(found) == 1, (pattern, [p.text for p in doc.paragraphs])
    return found[0]


def shading(el):
    shd = el.find(".//" + qn("w:shd"))
    return None if shd is None else shd.get(qn("w:fill")).upper()


def color(run):
    return str(run.font.color.rgb)


# ---------- 완료 기준 1: 번호·상자·목록·색·글꼴 ----------

def test_heading_numbers(d):
    assert para_match(d, r"01\s+첫 장").style.name == "Heading 1"
    assert para_match(d, r"02\s+둘째 장").style.name == "Heading 1"
    assert para_match(d, r"1\.1\s+첫 절").style.name == "Heading 2"
    assert para_match(d, r"1\.2\s+둘째 절").style.name == "Heading 2"
    assert para_match(d, r"2\.1\s+셋째 절").style.name == "Heading 2"
    assert para(d, "소절").style.name == "Heading 3"


def test_callout_is_blue_box(d):
    p = para(d, "첫 장 핵심 문장입니다.")
    ppr = p._p.pPr
    assert shading(ppr) == "EAF6FF"
    left = ppr.find(qn("w:pBdr")).find(qn("w:left"))
    assert left.get(qn("w:color")).upper() == "0099FF"


def test_list_two_levels(d):
    one = para_match(d, r"1\.\s+가")
    sub = para_match(d, r"\(1\)\s+가1")
    para_match(d, r"\(2\)\s+가2")
    para_match(d, r"2\.\s+나")
    bullet = para_match(d, r"•\s+하나")
    dash = para_match(d, r"–\s+둘")

    assert sub.paragraph_format.left_indent > (one.paragraph_format.left_indent or 0)
    assert dash.paragraph_format.left_indent > (bullet.paragraph_format.left_indent or 0)


def test_inline_marks_become_styles(d):
    p = para(d, "본문에 강조와 applies 코드가 있습니다.")
    strong = next(r for r in p.runs if r.text == "강조")
    code = next(r for r in p.runs if r.text == "applies")

    assert strong.bold and color(strong) == "0072C6"
    assert code.font.name == "Consolas"
    body_text = "\n".join(p.text for p in d.paragraphs)
    assert "**" not in body_text and "`" not in body_text


def test_evidence_is_small_gray(d):
    p = para(d, "근거: 40 CFR 70.6")
    assert all(r.font.size == Pt(9) and color(r) == "6B6B6B" for r in p.runs)


def test_table_colors(d):
    table = d.tables[0]
    header = table.rows[0].cells[0]
    assert shading(header._tc) == "F2F2F2"
    assert "D9D9D9" in table._tbl.xml.upper()

    adopted = table.cell(1, 1).paragraphs[0].runs[0]
    rejected = table.cell(2, 1).paragraphs[0].runs[0]
    assert color(adopted) == "15803D"
    assert color(rejected) == "B91C1C"


def test_qa_and_code_are_gray_boxes(d):
    q = para(d, "Q. 질문인가요?")
    a = para(d, "A. 답입니다.")
    assert shading(q._p.pPr) == "F2F2F2" and shading(a._p.pPr) == "F2F2F2"
    assert q.runs[0].bold and a.runs[0].bold

    for line in ("원문 줄", "  들여쓴 줄"):
        p = para(d, line)
        assert shading(p._p.pPr) == "F2F2F2"
        assert all(r.font.name == "Consolas" for r in p.runs)


def test_circled_number_is_bold(d):
    p = para(d, "① 첫째 포인트.")
    assert p.runs[0].text.startswith("①") and p.runs[0].bold


def test_styles_font_color_spacing(d):
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        fonts = d.styles[name].element.rPr.find(qn("w:rFonts"))
        assert fonts.get(qn("w:ascii")) == "Pretendard", name
        assert fonts.get(qn("w:eastAsia")) == "Pretendard", name

    normal = d.styles["Normal"]
    assert normal.font.size == Pt(10.5)
    assert str(normal.font.color.rgb) == "1A1A1A"
    assert normal.paragraph_format.space_after == Pt(6)

    spacing = {"Heading 1": (24, 12), "Heading 2": (18, 6), "Heading 3": (12, 4)}
    for name, (before, after) in spacing.items():
        pf = d.styles[name].paragraph_format
        assert (pf.space_before, pf.space_after) == (Pt(before), Pt(after)), name


# ---------- 완료 기준 2: 표지·목차·쪽 번호 ----------

def page_break_before(p):
    return bool(p.paragraph_format.page_break_before or p.style.paragraph_format.page_break_before)


def test_cover_is_own_page(d):
    assert d.paragraphs[0].text == "문서 제목"
    assert d.paragraphs[0].style.name == "Title"
    for p in d.paragraphs:
        if p.style.name == "Heading 1":
            assert page_break_before(p), p.text


def has_toc(doc):
    return "TOC \\o" in doc.element.body.xml


def test_toc_only_when_asked(d):
    assert not has_toc(d)

    with_toc = build(SAMPLE, toc=True)
    xml = with_toc.element.body.xml
    assert has_toc(with_toc)
    assert xml.index("TOC \\o") < xml.index("첫 장 핵심 문장")


def test_page_number_bottom_center(d):
    footer = d.sections[0].footer
    ps = [p for p in footer.paragraphs if "PAGE" in p._p.xml]
    assert ps and ps[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


# ---------- 완료 기준 3: 검사 오류면 Word를 만들지 않는다 ----------

def test_main_writes_docx_next_to_md(tmp_path):
    md = tmp_path / "a.md"
    md.write_text(SAMPLE, encoding="utf-8")

    assert main([str(md)], pdf=False) == 0
    assert has_toc(Document(tmp_path / "a.docx")) is False


def test_main_skips_docx_when_check_fails(tmp_path):
    md = tmp_path / "a.md"
    md.write_text(SAMPLE.replace("> 둘째 장 핵심 문장입니다.\n\n", ""), encoding="utf-8")

    assert main([str(md)], pdf=False) == 1
    assert not (tmp_path / "a.docx").exists()


def test_project_docs_build(tmp_path):
    names = ["0_Project_summary.md", "1_Project_full.md"]
    for name in names:
        shutil.copy(PROJECT / name, tmp_path / name)

    assert main([str(tmp_path / n) for n in names], pdf=False) == 0
    assert not has_toc(Document(tmp_path / "0_Project_summary.docx"))
    assert has_toc(Document(tmp_path / "1_Project_full.docx"))
